from __future__ import annotations

import logging
from pathlib import Path
from typing import Any


SUPPORTED_DOCUMENT_EXTENSIONS = {".txt", ".md", ".docx", ".pdf"}
LOGGER = logging.getLogger(__name__)


def load_text_file(path: str | Path, encoding: str = "utf-8") -> str:
    target = Path(path)
    encodings = [encoding, "utf-8-sig", "gbk", "gb18030"]
    last_error: UnicodeDecodeError | None = None
    for candidate in dict.fromkeys(encodings):
        try:
            return target.read_text(encoding=candidate)
        except UnicodeDecodeError as exc:
            last_error = exc
    if last_error is not None:
        raise last_error
    return target.read_text(encoding=encoding)


def load_document(file_path: str) -> dict[str, Any]:
    path = Path(file_path)
    if not path.exists():
        return _error(path, f"File does not exist: {path}")

    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_DOCUMENT_EXTENSIONS:
        return _error(
            path,
            "Only .txt, .md, .docx, and text-based .pdf files are supported.",
        )

    try:
        text = _load_content_by_type(path)
    except Exception as exc:
        return _error(path, str(exc))

    if not text.strip():
        return _error(path, "The document does not contain readable text.")

    lines = text.splitlines()
    LOGGER.info(
        "Document parsed: file_name=%s file_type=%s char_count=%s line_count=%s preview=%s",
        path.name,
        suffix.lstrip("."),
        len(text),
        len(lines),
        _preview(text),
    )
    return {
        "status": "success",
        "error_message": "",
        "file_path": str(path),
        "file_name": path.name,
        "file_type": suffix.lstrip("."),
        "text": text,
        "line_count": len(lines),
        "char_count": len(text),
    }


def _load_content_by_type(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md"}:
        return load_text_file(path)
    if suffix == ".docx":
        return _load_docx(path)
    if suffix == ".pdf":
        return _load_pdf(path)
    raise ValueError(f"Unsupported document type: {suffix}")


def _load_docx(path: Path) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise ImportError("python-docx is required to parse .docx files.") from exc

    document = Document(str(path))
    parts = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    parts.extend(_docx_table_text(document))
    return "\n".join(parts)


def _docx_table_text(document: Any) -> list[str]:
    table_texts: list[str] = []
    for table_index, table in enumerate(document.tables, start=1):
        rows = table.rows
        if not rows:
            continue
        headers = [_normalize_cell(cell.text) for cell in rows[0].cells]
        has_header = any(headers)
        data_rows = rows[1:] if has_header else rows
        for row_index, row in enumerate(data_rows, start=1):
            cells = [_normalize_cell(cell.text) for cell in row.cells]
            if not any(cells):
                continue
            if has_header:
                pairs = [
                    f"{headers[index] or f'列{index + 1}'}：{value}"
                    for index, value in enumerate(cells)
                    if value
                ]
            else:
                pairs = [
                    f"列{index + 1}：{value}"
                    for index, value in enumerate(cells)
                    if value
                ]
            if pairs:
                table_texts.append(f"表格{table_index} 第{row_index}行：" + "；".join(pairs))
    return table_texts


def _normalize_cell(text: str) -> str:
    return " ".join(str(text or "").split()).strip()


def _load_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise ImportError("pypdf is required to parse text-based .pdf files.") from exc

    reader = PdfReader(str(path))
    pages = [(page.extract_text() or "").strip() for page in reader.pages]
    text = "\n".join(page for page in pages if page)
    if not text.strip():
        raise ValueError("PDF text could not be extracted. Scanned PDF OCR is not supported.")
    return text


def _error(path: Path, message: str) -> dict[str, Any]:
    LOGGER.warning("Document parse failed: file_name=%s error=%s", path.name, message)
    return {
        "status": "error",
        "error_message": message,
        "file_path": str(path),
        "file_name": path.name,
        "file_type": path.suffix.lower().lstrip("."),
        "text": "",
        "line_count": 0,
        "char_count": 0,
    }


def _preview(text: str, limit: int = 300) -> str:
    return " ".join(text.split())[:limit]


def split_text(
    *,
    document_id: str | None,
    file_name: str,
    text: str,
    chunk_size: int = 800,
    chunk_overlap: int = 100,
) -> list[dict[str, Any]]:
    normalized = text.strip()
    if not normalized:
        return []

    safe_chunk_size = max(100, chunk_size)
    safe_overlap = min(max(0, chunk_overlap), safe_chunk_size - 1)
    chunks: list[dict[str, Any]] = []
    start = 0
    index = 1

    while start < len(normalized):
        end = min(start + safe_chunk_size, len(normalized))
        chunk_text = normalized[start:end].strip()
        if chunk_text:
            chunks.append(
                {
                    "document_id": document_id,
                    "chunk_id": f"chunk-{index}",
                    "file_name": file_name,
                    "text": chunk_text,
                    "start_index": start,
                    "end_index": end,
                    "score": 0.0,
                }
            )
            index += 1
        if end >= len(normalized):
            break
        start = max(end - safe_overlap, start + 1)

    if chunks:
        lengths = [len(str(chunk.get("text") or "")) for chunk in chunks]
        LOGGER.info(
            "Text split completed: document_id=%s file_name=%s char_count=%s chunk_size=%s chunk_overlap=%s chunk_count=%s avg_chunk_length=%s",
            document_id,
            file_name,
            len(normalized),
            safe_chunk_size,
            safe_overlap,
            len(chunks),
            round(sum(lengths) / max(len(lengths), 1), 2),
        )
        for chunk in chunks[:10]:
            LOGGER.info(
                "Text chunk created: document_id=%s chunk_id=%s start_index=%s end_index=%s text_length=%s",
                document_id,
                chunk.get("chunk_id"),
                chunk.get("start_index"),
                chunk.get("end_index"),
                len(str(chunk.get("text") or "")),
            )
    return chunks
